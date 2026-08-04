"""Extract plain text from uploaded documents of various formats.

Reuses the project's existing parser (skills/quiz-mastery) for Office/PDF
formats instead of duplicating extraction logic.
"""
import os
import sys
from pathlib import Path

MAX_CHARS = 60000

TEXT_LIKE = {
    ".txt", ".text", ".md", ".markdown", ".csv", ".tsv",
    ".json", ".log", ".yaml", ".yml", ".xml", ".sql",
}
OFFICE_LIKE = {".docx", ".pptx", ".pdf"}

_SKILL_SRC = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                          "skills", "quiz-mastery", "src")


def _skill_parse(path):
    """Delegate to the in-repo parser (stdlib for docx/pptx, pymupdf for pdf)."""
    if _SKILL_SRC not in sys.path:
        sys.path.insert(0, _SKILL_SRC)
    from quiz_mastery.file_parser import parse_file
    return parse_file(str(path))


def _parse_xlsx(path):
    from openpyxl import load_workbook
    wb = load_workbook(str(path), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"=== Лист: {ws.title} ===")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                out.append("\t".join(cells))
    return "\n".join(out)


def _read_text(path):
    raw = Path(path).read_bytes()
    for enc in ("utf-8", "cp1251", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    raise ValueError("не удалось определить кодировку файла")


def extract_text(path, original_name=""):
    """Return (text, truncated). Raises ValueError with a human-readable reason."""
    suffix = (Path(original_name).suffix or Path(path).suffix).lower()

    if suffix in TEXT_LIKE:
        text = _read_text(path)
    elif suffix == ".xlsx":
        text = _parse_xlsx(path)
    elif suffix in OFFICE_LIKE:
        text = _skill_parse(path)
    else:
        shown = suffix or "(без расширения)"
        raise ValueError(
            f"формат {shown} не поддерживается. Доступны: "
            "txt, md, csv, json, docx, pptx, pdf, xlsx"
        )

    text = (text or "").strip()
    if not text:
        raise ValueError("файл прочитан, но текста в нём не найдено")

    if len(text) > MAX_CHARS:
        return text[:MAX_CHARS], True
    return text, False
